"""Tests for the publication-artifact engine (services/publication/)."""

import pytest

from services.publication import (
    build_publication_items,
    categories_to_items,
    list_artifacts,
    render_artifact,
    SUPPORTED_FIGURE_FORMATS,
    SUPPORTED_TABLE_FORMATS,
)

PUBLICATION_CATEGORIES = ("agreement", "heatmaps", "performance", "statistical")


def build_publication_artifacts(results, mode, categories, figure_formats=("svg",),
                                table_formats=("csv",)):
    """Crate files of the categories' views at their defaults."""
    items = categories_to_items(categories, figure_formats, table_formats)
    return build_publication_items(results, mode, items)[0]

MEASURED = {
    "n_shots": 128,
    "seed": 12345,
    "divergence_report": {
        "simulators": ["qiskit", "qibo"],
        "statistical_distance_matrix (js_divergence)": [[-1.0, 0.69], [0.69, -1.0]],
        "hellinger_distance_matrix": [[-1.0, 1.0], [1.0, -1.0]],
        "pairwise_agreement": [
            {
                "simulators": ["qiskit", "qibo"],
                "mode": "measured",
                "agreement": "agree_after_bit_reversal",
                "registered_transform": "bit_reversal",
                "metrics": {
                    "js_divergence": 0.6931,
                    "statistical_test": {
                        "basis": "bit_reversal", "verdict": "consistent",
                        "effect_size": {"total_variation": 0.0, "cramers_v": 0.0},
                        "power": {"power_to_detect_delta_at_observed_shots": 0.126,
                                  "shots_per_backend_for_target_power": 1570},
                        "chi2": {"p_value": 1.0, "p_value_by": 1.0, "p_value_bh": 1.0},
                        "g_test": {"p_value": 1.0, "p_value_by": 1.0, "p_value_bh": 1.0},
                    },
                },
            }
        ],
        "statistical_test_summary": {"family_size": 1, "verdict_counts": {"consistent": 1}},
    },
    "performance_report": {"ranking": [
        {"simulator": "qiskit", "status": "success", "compilation_time_sec": 0.01,
         "simulation_time_sec": 0.08, "execution_time_sec": 0.09, "cpu_time_sec": 0.085,
         "gate_count": 3, "circuit_depth": 3},
        {"simulator": "qibo", "status": "success", "compilation_time_sec": 0.02,
         "simulation_time_sec": 0.03, "execution_time_sec": 0.05, "cpu_time_sec": 0.048,
         "gate_count": 3, "circuit_depth": 3},
    ]},
    "resource_report": {"ranking": [
        {"simulator": "qibo", "status": "success", "memory_usage_mb": 10.0, "process_peak_mb": 100.0,
         "process_peak_rss_mb": 95.5},
        {"simulator": "qiskit", "status": "success", "memory_usage_mb": 12.5, "process_peak_mb": 120.0,
         "process_peak_rss_mb": 110.3},
    ]},
    "raw_results": [
        {"simulator": "qiskit", "counts": {"001": 128}},
        {"simulator": "qibo", "counts": {"100": 128}},
    ],
}

STATEVECTOR = {
    "divergence_report": {
        "simulators": ["qiskit", "cirq"],
        "fidelity_matrix": [[-1.0, 0.0], [0.0, -1.0]],
        "pairwise_agreement": [
            {"simulators": ["qiskit", "cirq"], "mode": "statevector",
             "agreement": "agree_after_bit_reversal", "registered_transform": "bit_reversal",
             "metrics": {"direct_fidelity": 0.0, "bit_reversed_fidelity": 1.0}},
        ],
    },
    "performance_report": {"ranking": [
        {"simulator": "qiskit", "status": "success", "compilation_time_sec": 0.01,
         "simulation_time_sec": 0.04, "execution_time_sec": 0.05},
    ]},
    "resource_report": {"ranking": []},
}

_MAGIC = {
    "svg": lambda b: b"<svg" in b[:400],
    "png": lambda b: b[:8] == b"\x89PNG\r\n\x1a\n",
    "pdf": lambda b: b[:5] == b"%PDF-",
    "xlsx": lambda b: b[:2] == b"PK",
}


def _all_measured():
    return build_publication_artifacts(
        MEASURED, "measured", PUBLICATION_CATEGORIES,
        figure_formats=SUPPORTED_FIGURE_FORMATS, table_formats=SUPPORTED_TABLE_FORMATS)


def test_all_figure_formats_have_valid_magic_bytes():
    arts = _all_measured()
    figs = {p: b for p, b in arts.items() if p.startswith("figures/") and not p.endswith(".json")}
    assert figs
    for path, data in figs.items():
        ext = path.rsplit(".", 1)[1]
        assert _MAGIC[ext](data), f"{path} failed magic check"


def test_all_tables_present_in_all_formats():
    arts = _all_measured()
    for base in ("agreement", "execution_time", "resource_usage", "statistical_test"):
        for ext in ("csv", "tex", "html", "xlsx"):
            assert f"tables/{base}.{ext}" in arts, f"missing tables/{base}.{ext}"
    assert _MAGIC["xlsx"](arts["tables/agreement.xlsx"])


def test_artifacts_are_byte_deterministic():
    """Every non-XLSX artifact is byte-identical across two builds."""
    a1 = _all_measured()
    a2 = _all_measured()
    assert a1.keys() == a2.keys()
    for path in a1:
        if path.endswith(".xlsx"):
            continue
        assert a1[path] == a2[path], f"{path} not reproducible"


def test_xlsx_tables_are_valid_and_reopenable():
    import io
    from openpyxl import load_workbook
    arts = _all_measured()
    xlsx = [b for p, b in arts.items() if p.endswith(".xlsx")]
    assert xlsx
    for data in xlsx:
        assert data[:2] == b"PK"
        wb = load_workbook(io.BytesIO(data))
        assert wb.active.max_row >= 3


def test_latex_uses_booktabs_and_escapes_underscores():
    arts = build_publication_artifacts(MEASURED, "measured", ("agreement",),
                                       (), ("latex",))
    tex = arts["tables/agreement.tex"].decode()
    assert "\\toprule" in tex and "\\bottomrule" in tex and "booktabs" in tex
    assert "agree\\_after\\_bit\\_reversal" in tex
    assert "registered-transform lattice" in tex


def test_agreement_table_content():
    arts = build_publication_artifacts(MEASURED, "measured", ("agreement",), (), ("csv",))
    csv = arts["tables/agreement.csv"].decode()
    assert "qiskit × qibo" in csv
    assert "agree_after_bit_reversal" in csv
    assert "consistent" in csv


def test_statistical_test_table_measured_only():
    measured = build_publication_artifacts(MEASURED, "measured", ("statistical",), (), ("csv",))
    assert "tables/statistical_test.csv" in measured
    csv = measured["tables/statistical_test.csv"].decode()
    assert "1570" in csv and "consistent" in csv
    sv = build_publication_artifacts(STATEVECTOR, "statevector", ("statistical",),
                                     SUPPORTED_FIGURE_FORMATS, SUPPORTED_TABLE_FORMATS)
    assert not any("statistical_test" in p for p in sv)
    assert not any("histogram" in p for p in sv)


def test_mode_branching_heatmaps():
    sv = build_publication_artifacts(STATEVECTOR, "statevector", ("heatmaps",),
                                     ("svg",), ())
    assert "figures/fidelity_heatmap.svg" in sv
    assert not any("js_divergence" in p or "hellinger" in p for p in sv)
    measured = build_publication_artifacts(MEASURED, "measured", ("heatmaps",), ("svg",), ())
    assert "figures/js_divergence_heatmap.svg" in measured
    assert "figures/hellinger_heatmap.svg" in measured
    assert not any("fidelity" in p for p in measured)


def test_heatmap_handles_sentinel_and_errored_runner():
    report = {
        "divergence_report": {
            "simulators": ["a", "b", "c"],
            "fidelity_matrix": [[-1.0, 1.0, -1.0], [1.0, -1.0, -1.0], [-1.0, -1.0, -1.0]],
            "pairwise_agreement": [],
        },
    }
    arts = build_publication_artifacts(report, "statevector", ("heatmaps",), ("svg",), ())
    assert "figures/fidelity_heatmap.svg" in arts
    assert _MAGIC["svg"](arts["figures/fidelity_heatmap.svg"])


def test_graceful_when_reports_absent():
    assert build_publication_artifacts({}, "measured", PUBLICATION_CATEGORIES,
                                       SUPPORTED_FIGURE_FORMATS, SUPPORTED_TABLE_FORMATS) == {}
    no_pairwise = {"divergence_report": {"simulators": ["a", "b"], "fidelity_matrix": [[-1.0, 1.0], [1.0, -1.0]]}}
    arts = build_publication_artifacts(no_pairwise, "statevector", ("agreement", "heatmaps"), ("svg",), ("csv",))
    assert not any("agreement" in p for p in arts)
    assert "figures/fidelity_heatmap.svg" in arts


def test_unsupported_formats_ignored():
    arts = build_publication_artifacts(MEASURED, "measured", ("agreement",),
                                       ("svg", "bogus"), ("csv", "nope"))
    assert "figures/agreement_distribution.svg" in arts
    assert "tables/agreement.csv" in arts
    assert not any(p.endswith(".bogus") or p.endswith(".nope") for p in arts)

def test_catalog_lists_expected_artifacts_measured():
    ids = {a["id"] for a in list_artifacts(MEASURED, "measured")}
    assert {"agreement_distribution", "agreement", "js_divergence_heatmap",
            "hellinger_heatmap", "performance_time", "resource_memory",
            "performance", "resources", "statistical_test"} <= ids
    assert {"matrix_js", "matrix_hellinger", "metadata", "glossary"} <= ids
    assert {"counts_overview", "counts_table", "qubit_marginals", "qubit_marginals_table",
            "hamming_weight", "hamming_weight_table", "rank_probability",
            "rank_probability_table"} <= ids
    assert not any(i.startswith("histogram_") for i in ids)
    assert "fidelity_heatmap" not in ids


def test_catalog_statevector_branches():
    ids = {a["id"] for a in list_artifacts(STATEVECTOR, "statevector")}
    assert "fidelity_heatmap" in ids and "matrix_fidelity" in ids
    assert not any("hellinger" in i or "js_divergence" in i for i in ids)
    assert not ids & {"statistical_test", "counts_overview", "qubit_marginals", "hamming_weight",
                      "rank_probability"}
    assert "glossary" in ids


def test_catalog_entries_have_view_formats_legend():
    for a in list_artifacts(MEASURED, "measured"):
        assert a["view"] in ("image", "text")
        assert a["formats"] and a["legend"]
        if a["view"] == "image":
            assert set(a["formats"]) == set(SUPPORTED_FIGURE_FORMATS)


def test_render_artifact_image_scale_increases_size():
    small, mime, legend = render_artifact(MEASURED, "measured", "agreement_distribution", "png", 1)
    big, _, _ = render_artifact(MEASURED, "measured", "agreement_distribution", "png", 5)
    assert small[:8] == b"\x89PNG\r\n\x1a\n" and big[:8] == b"\x89PNG\r\n\x1a\n"
    assert len(big) > len(small)
    assert mime == "image/png" and legend


def test_render_artifact_text_formats():
    for fmt in ("latex", "csv", "html", "markdown"):
        data, mime, _ = render_artifact(MEASURED, "measured", "agreement", fmt)
        assert data and isinstance(data, bytes)
    md = render_artifact(MEASURED, "measured", "agreement", "markdown")[0].decode()
    assert md.startswith("**") and "|" in md


def test_render_new_artifacts_content():
    mjs = render_artifact(MEASURED, "measured", "matrix_js", "csv")[0].decode()
    assert "qiskit" in mjs and "qibo" in mjs
    meta = render_artifact(MEASURED, "measured", "metadata", "csv")[0].decode()
    assert "qiskit" in meta
    gloss = render_artifact(MEASURED, "measured", "glossary", "markdown")[0].decode()
    assert "agree_after_bit_reversal" in gloss and "consistent" in gloss


def test_render_artifact_rejects_bad_id_and_format():
    with pytest.raises(KeyError):
        render_artifact(MEASURED, "measured", "does_not_exist", "png")
    with pytest.raises(ValueError):
        render_artifact(MEASURED, "measured", "agreement", "png")


def test_counts_overview_available_measured_only():
    ids = {a["id"] for a in list_artifacts(MEASURED, "measured")}
    assert "counts_overview" in ids
    assert "counts_overview" not in {a["id"] for a in list_artifacts(STATEVECTOR, "statevector")}


def test_counts_overview_renders_all_figure_formats():
    for fmt in SUPPORTED_FIGURE_FORMATS:
        data, mime, legend = render_artifact(MEASURED, "measured", "counts_overview", fmt)
        assert _MAGIC[fmt](data) and legend


def test_count_views_need_a_backend_with_counts():
    one = dict(MEASURED, raw_results=[{"simulator": "qiskit", "counts": {"001": 128}}])
    assert "counts_overview" in {a["id"] for a in list_artifacts(one, "measured")}
    none = dict(MEASURED, raw_results=[{"simulator": "qiskit", "error": "boom"}])
    assert "counts_overview" not in {a["id"] for a in list_artifacts(none, "measured")}


def test_counts_overview_in_statistical_bundle():
    files = build_publication_artifacts(MEASURED, "measured", ("statistical",),
                                        figure_formats=("svg",), table_formats=("csv",))
    assert "figures/counts.svg" in files


def test_counts_heatmap_layout_renders():
    data, _, legend = render_artifact(MEASURED, "measured", "counts_overview", "png",
                                      params={"layout": "heatmap"})
    assert _MAGIC["png"](data) and legend.startswith("Outcome × backend heatmap")


def test_count_comparison_shows_the_chosen_number_of_outcomes():
    import matplotlib.pyplot as plt
    from services.publication.figures_counts import fig_counts
    from services.publication.views import normalize_params
    sims = ["qiskit", "qibo", "cirq", "pennylane", "braket"]
    rep = {"n_shots": 1756, "raw_results": [
        {"simulator": s, "counts": {**{format(i, "08b"): 1 for i in range(256)}, "00000000": 1500}}
        for s in sims]}
    p = normalize_params("counts", {"top": 10}, rep)
    fig = fig_counts(rep, p)
    try:
        assert len(fig.axes[0].get_xticklabels()) == 11
    finally:
        plt.close(fig)


def _report_with_failures():
    from tests.unit.test_artifacts_parity import _sv_edge
    return _sv_edge()


def test_tables_list_unranked_runs_after_ranked_with_reason():
    rep = _report_with_failures()
    perf = render_artifact(rep, "statevector", "performance", "csv")[0].decode().splitlines()
    assert [line.split(",")[1] for line in perf[1:]] == ["s0", "s1", "s2", "s3"]
    assert perf[3] == "—,s2,CPU,—,—,—,—,—,—,—,timeout"
    assert perf[4].endswith("error (missing_timing)")
    res = render_artifact(rep, "statevector", "resources", "csv")[0].decode().splitlines()
    assert any(line.endswith("ok (missing_process_peak_rss)") for line in res)
    assert "—,s2,CPU,—,—,—,—,—,timeout" in res


def test_ranking_figures_mark_unranked_runs_and_plot_the_ranking_metric():
    import matplotlib.pyplot as plt
    from services.publication import _fig_ranking_bar
    rep = _report_with_failures()
    fig = _fig_ranking_bar(rep, "resource_report",
                           [("process_peak_rss_mb", "#2a78d6", "process peak RSS")], "t", "MiB")
    try:
        ax = fig.axes[0]
        widths = [p.get_width() for p in ax.patches]
        assert widths == [80.0, 0.0, 0.0]
        notes = [t.get_text() for t in ax.texts]
        assert any("missing_process_peak_rss" in n for n in notes)
        assert any(n.strip() == "not ranked: timeout" for n in notes)
    finally:
        plt.close(fig)
    only_failed = {"performance_report": {"ranking": [], "unranked": [
        {"simulator": "x", "status": "error", "reason": "runner_error"}]}}
    data, _, legend = render_artifact(only_failed, "statevector", "performance_time", "svg")
    assert _MAGIC["svg"](data) and "n=1 timing sample" in legend


def test_heatmap_annotations_do_not_round_to_one_or_zero():
    import matplotlib.pyplot as plt
    from services.publication import _fig_heatmap
    rep = _report_with_failures()
    fig = _fig_heatmap(rep["divergence_report"], "fidelity_matrix", "t")
    try:
        texts = sorted(t.get_text() for t in fig.axes[0].texts)
    finally:
        plt.close(fig)
    assert "0.99996" in texts and "0.996" in texts and "5.00e-05" in texts
    assert texts.count("1.00") == 2


def test_fmt_num_never_shows_non_one_as_one_or_nonzero_as_zero():
    from services.publication import _fmt_num
    assert _fmt_num(0.99996) == "0.99996"
    assert _fmt_num(0.996, 2) == "0.996"
    assert _fmt_num(4e-5) == "4.00e-05"
    assert _fmt_num(6.103515625e-05) == "6.10e-05"
    assert _fmt_num(0) == "0.0000" and _fmt_num(1) == "1.0000"


def test_ranking_figures_state_timing_scope_in_subtitle():
    rep = _report_with_failures()
    perf = render_artifact(rep, "statevector", "performance_time", "svg")[0].decode()
    assert "n=1 timing sample(s) per backend; timed run only, warm-up excluded" in perf
    mem = render_artifact(rep, "statevector", "resource_memory", "svg")[0].decode()
    assert "SDK import dominates for small circuits" in mem


def test_heatmap_masks_null_cells():
    rep = {"divergence_report": {"simulators": ["a", "b", "c"],
                                 "fidelity_matrix": [[1.0, 0.5, None], [0.5, 1.0, None],
                                                     [None, None, None]]}}
    data = render_artifact(rep, "statevector", "fidelity_heatmap", "svg")[0]
    assert _MAGIC["svg"](data)
    table = render_artifact(rep, "statevector", "matrix_fidelity", "csv")[0].decode().splitlines()
    assert table[1] == "a,—,0.5,—"


def test_count_figures_use_common_msb_key_order_and_name_the_run():
    from services.publication.selectors import _msb_counts, _run_label
    results = {"n_shots": 100, "task_name": "x (3q)", "raw_results": [
        {"simulator": "q", "qubit_ordering": "lsb", "counts": {"001": 100}},
        {"simulator": "c", "qubit_ordering": "msb", "counts": {"100": 100}},
        {"simulator": "u", "counts": {"001": 100}},
        {"simulator": "e", "error": "boom"},
    ]}
    assert _msb_counts(results) == {"q": {"100": 100}, "c": {"100": 100}, "u": {"001": 100}}
    assert _run_label(results) == "x (3q), 100 shots per backend"
    assert _run_label({}) == "unnamed circuit"


def test_table_columns_use_one_notation_and_csv_keeps_full_precision():
    from services.publication import _tbl_performance, _table_bytes
    rep = {"performance_report": {"ranking": [
        {"simulator": "a", "rank": 1, "execution_time_sec": 0.000825, "compilation_time_sec": 0.0001,
         "simulation_time_sec": 0.000725},
        {"simulator": "b", "rank": 2, "execution_time_sec": 0.001331123456789,
         "compilation_time_sec": 0.0003, "simulation_time_sec": 0.001031123456789},
    ], "unranked": []}, "raw_results": []}
    headers, rows, title, caption = _tbl_performance(rep)
    total = headers.index("Total (s)")
    assert [r[total] for r in rows] == ["8.25e-04", "1.33e-03"]
    csv = _table_bytes(headers, rows, title, caption, "csv").decode().splitlines()
    assert csv[2].split(",")[total] == "0.001331123456789"
    assert "GPU runners execute kernels on the GPU" in caption


def test_figures_name_the_circuit_and_mark_gpu_runs():
    import matplotlib.pyplot as plt
    from services.publication import _fig_ranking_bar
    from services.publication.selectors import _run_label
    rep = {"circuit_name": "GHZ (3q)", "n_shots": 10, "performance_report": {"ranking": [
        {"simulator": "cu", "device": "GPU", "rank": 1, "execution_time_sec": 0.1,
         "compilation_time_sec": 0.0, "simulation_time_sec": 0.1}], "unranked": []},
        "raw_results": []}
    assert _run_label(rep).startswith("GHZ (3q)")
    svg = render_artifact(rep, "measured", "performance_time", "svg")[0].decode()
    assert "GHZ (3q)" in svg
    fig = _fig_ranking_bar(rep, "performance_report",
                           [("execution_time_sec", "#2a78d6", "total")], "t", "s")
    try:
        assert [t.get_text() for t in fig.axes[0].get_yticklabels()] == ["cu (GPU)"]
    finally:
        plt.close(fig)


@pytest.mark.parametrize("name", ["a $x^$ b", "GHZ $\\alpha$ test", "$\\notacommand{x}$"])
def test_user_names_with_dollar_signs_render_literally(name):
    import copy
    rep = copy.deepcopy(MEASURED)
    rep["run_name"] = name
    for aid, params in (("performance_time", {}), ("resource_memory", {}), ("counts_overview", {}),
                        ("counts_overview", {"layout": "heatmap"}),
                        ("counts_overview", {"layout": "panels"}), ("qubit_marginals", {}),
                        ("hamming_weight", {}), ("rank_probability", {})):
        svg = render_artifact(rep, "measured", aid, "svg", params=params)[0].decode()
        assert name in svg.replace("<!--", ""), aid
        assert "α" not in svg, aid
    files = build_publication_artifacts(rep, "measured", PUBLICATION_CATEGORIES,
                                        figure_formats=("svg", "png", "pdf"))
    assert name in files["figures/execution_time.svg"].decode()


def _repeated(results):
    import copy
    rep = copy.deepcopy(results)
    rep["raw_results"] = [
        {"simulator": "qiskit", "counts": {"001": 128}, "timing_samples": 5,
         "compilation_time_samples": [0.01, 0.012, 0.011, 0.02, 0.01],
         "simulation_time_samples": [0.08, 0.07, 0.3, 0.09, 0.075],
         "compilation_time_sec": 0.011, "simulation_time_sec": 0.08, "execution_time_sec": 0.091},
        {"simulator": "qibo", "counts": {"100": 128}, "timing_samples": 5,
         "compilation_time_samples": [0.02, 0.02, 0.021, 0.019, 0.02],
         "simulation_time_samples": [0.03, 0.031, 0.029, 0.03, 0.2],
         "compilation_time_sec": 0.02, "simulation_time_sec": 0.03, "execution_time_sec": 0.05},
    ]
    rep["timing_samples"] = 5
    return rep


def test_repeated_timing_table_shows_median_iqr_stdev_and_n():
    import csv, io, statistics
    rep = _repeated(MEASURED)
    text = render_artifact(rep, "measured", "performance", "csv")[0].decode()
    rows = list(csv.reader(io.StringIO(text)))
    assert rows[0][3:9] == ["Compilation median (s)", "Simulation median (s)", "Total median (s)",
                            "Total IQR (s)", "Total stdev (s)", "n"]
    qiskit = next(r for r in rows if r[1] == "qiskit")
    totals = [c + s for c, s in zip(rep["raw_results"][0]["compilation_time_samples"],
                                    rep["raw_results"][0]["simulation_time_samples"])]
    assert float(qiskit[5]) == pytest.approx(0.011 + 0.08)
    assert float(qiskit[7]) == pytest.approx(statistics.stdev(totals))
    assert qiskit[8] == "5"
    caption = render_artifact(rep, "measured", "performance", "markdown")[0].decode()
    assert "Median wall-clock times" in caption and "Ranked by median total time" in caption


def test_performance_report_ranks_by_median_total_time():
    from services.analysis.performance import create_performance_report
    fast_median = {"simulator": "a", "execution_time_sec": 0.3, "compilation_time_sec": 0.0,
                   "simulation_time_sec": 0.3, "compilation_time_samples": [0.0, 0.0, 0.0],
                   "simulation_time_samples": [0.1, 0.1, 5.0], "timing_samples": 3}
    steady = {"simulator": "b", "execution_time_sec": 0.2, "compilation_time_sec": 0.0,
              "simulation_time_sec": 0.2, "compilation_time_samples": [0.0, 0.0, 0.0],
              "simulation_time_samples": [0.2, 0.2, 0.2], "timing_samples": 3}
    report = create_performance_report([fast_median, steady])
    assert report["ranking_metric"] == "execution_time_sec_median"
    assert [e["simulator"] for e in report["ranking"]] == ["a", "b"]
    first = report["ranking"][0]
    assert first["execution_time_sec_median"] == pytest.approx(0.1)
    assert first["timing_samples"] == 3 and first["timing_statistic"] == "median"
    single = create_performance_report([{"simulator": "c", "execution_time_sec": 0.1}])
    assert single["ranking_metric"] == "execution_time_sec"
    assert "execution_time_sec_median" not in single["ranking"][0]
